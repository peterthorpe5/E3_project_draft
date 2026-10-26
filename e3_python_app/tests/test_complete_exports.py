"""Lossless complete export regressions, including Excel limits and literal cells."""

from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

import duckdb
from openpyxl import load_workbook
import pytest

from e3app import complete_exports
from e3app.complete_exports import (
    _write_excel_row,
    complete_excel_workbook,
    complete_tsv_archive,
    iter_query_batches,
    query_schema,
)
from e3app.errors import AppError
from e3app.result_queries import ResourceQuery, build_result_query


def test_complete_tsv_contains_all_rows_all_fields_and_counts(resource_db: Path) -> None:
    """An arbitrary screen-sized dataset cannot truncate the archive."""
    with duckdb.connect(str(resource_db)) as connection:
        connection.execute(
            "CREATE TABLE many_rows AS SELECT i AS rank, 'N0.HOG' || i AS hog_id, "
            "repeat('x', 100) AS hidden_field FROM range(1507) t(i)"
        )
        query = build_result_query(connection=connection, result="many_rows")
        payload = complete_tsv_archive(
            connection=connection,
            queries=(query,),
            description="all rows, all fields",
        )
    with ZipFile(BytesIO(payload)) as archive:
        manifest = list(
            csv.DictReader(StringIO(archive.read("manifest.tsv").decode()), delimiter="\t")
        )
        assert manifest[0]["exported_rows"] == "1507"
        table = list(
            csv.DictReader(
                StringIO(archive.read(manifest[0]["file"]).decode()),
                delimiter="\t",
            )
        )
        assert len(table) == 1507
        assert table[-1]["rank"] == "1506"
        assert table[-1]["hidden_field"] == "x" * 100
        assert b"not fetched or embedded" in archive.read("README.txt")
        assert b"hidden_field" in archive.read("column_definitions.tsv")


def test_excel_splits_all_rows_and_keeps_formulas_as_literal_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Worksheet row limits cause splitting rather than data loss."""
    monkeypatch.setattr(complete_exports, "EXCEL_DATA_ROWS", 2)
    with duckdb.connect(":memory:") as connection:
        connection.execute(
            "CREATE TABLE source AS SELECT i AS final_rank, '=SUM(A1:A2)' AS protein_name, "
            "'https://example.org/protein' AS url FROM range(5) t(i)"
        )
        query = build_result_query(connection=connection, result="source")
        empty = ResourceQuery(name="empty", sql="SELECT 1 AS value WHERE FALSE")
        payload = complete_excel_workbook(
            connection=connection,
            queries=(query, empty),
            description="every record",
        )
    workbook = load_workbook(BytesIO(payload))
    assert workbook["Manifest"]["C2"].value == 5
    assert workbook["Manifest"]["C3"].value == 0
    assert workbook["Result_001_1"]["B2"].data_type == "s"
    assert workbook["Result_001_1"]["B2"].value == "=SUM(A1:A2)"
    assert workbook["Result_001_1"]["C2"].hyperlink is None
    assert [workbook[f"Result_001_{part}"]["A2"].value for part in (1, 2, 3)] == [0, 2, 4]
    assert workbook["Result_002_1"]["A1"].value == "value"
    assert workbook["Column definitions"].max_row == 5


def test_excel_rejects_long_cells_and_tsv_retains_every_character() -> None:
    """Excel's cell text limit must not silently truncate annotation or sequences."""
    long_value = "a\tb\n" + "X" * 40_000
    with duckdb.connect(":memory:") as connection:
        connection.execute("CREATE TABLE source(sequence VARCHAR)")
        connection.execute("INSERT INTO source VALUES (?)", [long_value])
        query = build_result_query(connection=connection, result="source")
        with pytest.raises(AppError, match="cell text limit"):
            complete_excel_workbook(
                connection=connection, queries=(query,), description="long cells"
            )
        payload = complete_tsv_archive(
            connection=connection, queries=(query,), description="long cells"
        )
    with ZipFile(BytesIO(payload)) as archive:
        table = list(
            csv.reader(StringIO(archive.read("tables/001_source.tsv").decode()), delimiter="\t")
        )
        assert table[1][0] == long_value


def test_batch_validation_and_failed_sources_return_no_partial_export() -> None:
    """Invalid batches and unreadable sources fail with actionable errors."""
    with duckdb.connect(":memory:") as connection:
        query = ResourceQuery(name="source", sql="SELECT * FROM missing")
        for batch_rows in (0, True, 1.5, 100001):
            with pytest.raises(AppError, match="batch rows"):
                list(iter_query_batches(connection=connection, query=query, batch_rows=batch_rows))
        with pytest.raises(AppError, match="Complete export failed"):
            list(iter_query_batches(connection=connection, query=query))
        with pytest.raises(AppError, match="inspect export fields"):
            query_schema(connection=connection, query=query)
        for exporter in (complete_tsv_archive, complete_excel_workbook):
            for queries, description in (
                ((), "scope"),
                ((query, query), "scope"),
                ((ResourceQuery(name=" ", sql="SELECT 1"),), "scope"),
                ((query,), ""),
            ):
                with pytest.raises(AppError):
                    exporter(connection=connection, queries=queries, description=description)


def test_excel_scalar_dates_nulls_and_limit_status() -> None:
    """Dates retain values and failed cell writes cannot escape detection."""
    output = BytesIO()
    import xlsxwriter

    with xlsxwriter.Workbook(output) as workbook:
        worksheet = workbook.add_worksheet()
        _write_excel_row(
            worksheet=worksheet,
            row_number=0,
            values=(date(2026, 10, 26), datetime(2026, 10, 26, 11), None, True, 7),
        )
        with pytest.raises(AppError, match="complete cell"):
            _write_excel_row(worksheet=worksheet, row_number=1_048_576, values=(1,))
        with pytest.raises(AppError, match="many fields"):
            _write_excel_row(worksheet=worksheet, row_number=1, values=(1,) * 16_385)
    workbook = load_workbook(BytesIO(output.getvalue()))
    assert workbook.active["A1"].value == "2026-10-26"
    assert workbook.active["B1"].value == "2026-10-26T11:00:00"


def test_complete_export_passes_the_old_hundred_thousand_row_ceiling() -> None:
    """Complete means every row even above the former query-layer maximum."""
    with duckdb.connect(":memory:") as connection:
        query = ResourceQuery(name="large", sql="SELECT i AS row_id FROM range(100007) t(i)")
        payload = complete_tsv_archive(
            connection=connection,
            queries=(query,),
            description="all 100007 rows",
        )
    with ZipFile(BytesIO(payload)) as archive:
        rows = archive.read("tables/001_large.tsv").decode().splitlines()
        assert len(rows) == 100008
        assert rows[-1] == "100006"
        assert "100007" in archive.read("manifest.tsv").decode()
