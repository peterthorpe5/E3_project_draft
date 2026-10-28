"""Bulk archive security, exact mappings and portable output tests."""

import io
import json
import stat
import zipfile

import duckdb
import pytest
from openpyxl import load_workbook

from pocket_workbench.core import (
    WorkbenchError,
    digest_file,
    inside,
    load_project,
    safe_id,
    utc_now,
    write_json,
    write_tsv,
)
from pocket_workbench.database import (
    export_tables,
    pack_project,
    query_table,
    save_project,
    workbook_bytes,
)
from pocket_workbench.help_content import field_help, glossary_rows, methods_text
from pocket_workbench.importing import (
    extract_zip,
    import_project,
    map_sequence,
    read_fasta,
)
from pocket_workbench.viewers import (
    alignment_document,
    json_for_html,
    sequence_position_columns,
    structure_document,
    viewer_asset,
)


@pytest.mark.parametrize(
    "filename", ["../escape.pdb", "/absolute.pdb", "C:/absolute.pdb", "a\\b.pdb"]
)
def test_zip_traversal_is_rejected(tmp_path, filename):
    """ZIP members cannot write outside the extraction root on either platform."""
    archive = tmp_path / "malicious.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr(filename, "data")
    with pytest.raises(WorkbenchError):
        extract_zip(archive=archive, destination=tmp_path / "unpack")


def test_zip_symlinks_and_case_duplicates_are_rejected(tmp_path):
    """Symlink entries and ambiguous duplicate paths fail before import."""
    for kind in ["symlink", "duplicate"]:
        archive = tmp_path / f"{kind}.zip"
        with zipfile.ZipFile(archive, "w") as handle:
            if kind == "symlink":
                entry = zipfile.ZipInfo("model.pdb")
                entry.external_attr = (stat.S_IFLNK | 0o777) << 16
                handle.writestr(entry, "../elsewhere")
            else:
                handle.writestr("Model.pdb", "data")
                handle.writestr("model.pdb", "other")
        with pytest.raises(WorkbenchError):
            extract_zip(archive=archive, destination=tmp_path / kind)


def test_archive_skips_msa_templates_and_full_data(tmp_path):
    """Bulk AF3 imports retain models without unpacking unnecessary large data."""
    archive = tmp_path / "models.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        for filename in [
            "fold_a/model.cif",
            "fold_a/templates/template.cif",
            "fold_a/msas/x.fasta",
            "fold_a/fold_a_full_data_0.json",
        ]:
            handle.writestr(filename, "fixture")
    result = extract_zip(archive=archive, destination=tmp_path / "unpack")
    assert [path.name for path in result] == ["model.cif"]


def test_exact_mapping_does_not_infer_a_sequence_from_the_model_name():
    """Only exact content or an audited explicit fragment mapping is accepted."""
    assert map_sequence(sequence="ACDE", sequences={"one": "ACDE"}) == (
        "one",
        "exact_sequence",
        0,
    )
    assert (
        map_sequence(sequence="ACDE", sequences={"one": "ACDE", "two": "ACDE"})[0] == ""
    )
    assert map_sequence(sequence="ACDE", sequences={"one": "XXACDEXX"})[0] == ""
    assert map_sequence(
        sequence="ACDE", sequences={"one": "XXACDEXX"}, explicit_id="one"
    ) == ("one", "unique_exact_fragment_explicit", 2)
    with pytest.raises(WorkbenchError):
        map_sequence(sequence="AAAA", sequences={"one": "ACDE"}, explicit_id="one")


@pytest.mark.parametrize(
    "text", [">a\nACDE\n>a\nACDE\n", "ACDE\n", ">a\nA--C\n", ">\nACDE\n", ">a\n123\n"]
)
def test_fasta_rejects_ambiguous_or_invalid_records(tmp_path, text):
    """Duplicate headers, invalid residues and unexpected gaps fail explicitly."""
    path = tmp_path / "bad.fasta"
    path.write_text(text)
    with pytest.raises(WorkbenchError):
        read_fasta(path=path)


def test_native_project_zip_roundtrip_preserves_models_and_evidence(project, tmp_path):
    """A downloaded project can be reopened with typed evidence and local assets."""
    archive = pack_project(root=project, output=tmp_path / "portable.zip")
    restored = import_project(sources=[archive], output=tmp_path / "restored")
    assert len(load_project(root=restored)["models"]) == 3
    assert len(query_table(root=restored, table="pocket_matches")) == 4
    assert (restored / "models/synthetic_reference/prepared.pdb").is_file()


def test_report_import_without_coordinates_is_explicit(tmp_path):
    """A report-only project exposes tables without inventing model coordinates."""
    report = tmp_path / "report.html"
    report.write_text(
        '<script id="data" type="application/json">'
        + json.dumps(
            {"samples": [{"sample_id": "model", "status": "complete"}], "pockets": []}
        )
        + "</script><script>throw Error('never execute')</script>"
    )
    root = import_project(sources=[report], output=tmp_path / "resource")
    model = load_project(root=root)["models"][0]
    assert model["coordinates"] is None
    assert model["mapping_status"] == "coordinates_missing"
    assert query_table(root=root, table="residues") == []


def test_stale_project_write_is_rejected(project):
    """Two task snapshots cannot silently overwrite each other's results."""
    old = load_project(root=project)
    save_project(root=project, data=load_project(root=project))
    with pytest.raises(WorkbenchError, match="changed"):
        save_project(root=project, data=old)


def test_native_project_restore_keeps_prediction_reports_as_artifacts(
    project, tmp_path
):
    """Nested prediction reports must not override the portable app resource."""
    for identifier in ["job_one", "job_two"]:
        path = project / "prediction_runs" / identifier / "report.html"
        path.parent.mkdir(parents=True)
        path.write_text("This is an artifact, not the workbench's import source")
    archive = pack_project(root=project, output=tmp_path / "with_predictions.zip")
    restored = import_project(
        sources=[archive], output=tmp_path / "restored_predictions"
    )
    assert len(load_project(root=restored)["pockets"]) == 6
    assert (restored / "prediction_runs/job_one/report.html").is_file()


def test_excel_preserves_numbers_and_blocks_formula_injection():
    """Numeric metrics stay numeric and user-provided text remains literal."""
    value = workbook_bytes(
        tables={
            "Evidence": [
                {
                    "sample_id": '=HYPERLINK("bad")',
                    "symmetric_overlap": 0.8,
                    "centroid_distance": 1.2,
                }
            ]
        },
        notes="Methods",
    )
    book = load_workbook(io.BytesIO(value))
    sheet = book["Evidence"]
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 0.8
    assert sheet["B2"].number_format == "0.0%"
    assert sheet.freeze_panes == "C2"
    assert sheet.tables


def test_exports_are_tsv_and_database_has_no_external_views(project, tmp_path):
    """App integration uses materialised tables and TSV rather than CSV files."""
    output = tmp_path / "exports"
    export_tables(root=project, output=output)
    assert (output / "pocket_workbench.xlsx").is_file()
    assert (output / "pocket_matches.tsv").read_text().startswith("comparison_id\t")
    assert not list(output.glob("*.csv"))
    with duckdb.connect(str(project / "pocket_workbench.duckdb"), read_only=True) as db:
        assert db.execute("SELECT count(*) FROM residues").fetchone()[0] == 96
        assert (
            db.execute("SELECT count(*) FROM v_supported_pocket_matches").fetchone()[0]
            == 2
        )


def test_viewer_data_escapes_script_endings_and_uses_bundled_assets():
    """Hostile identifiers cannot terminate the inert data script."""
    encoded = json_for_html(value={"id": "</script><script>bad()</script>"})
    assert "</script>" not in encoded
    document = alignment_document(msa=">a\nACDE\n>b\nACDF\n", viewer="alignment")
    assert "AlignmentViewer" in document
    assert '<script src="https:' not in document
    overlay = structure_document(
        reference_pdb="END\n",
        mobile_pdb="END\n",
        reference_label="<hostile>",
        mobile_label="b",
    )
    assert "<hostile>" not in overlay
    assert "createViewer" in overlay
    assert "lining residues, not empty cavity volume" in overlay
    assert sequence_position_columns(aligned="A--C.D") == {1: 1, 2: 4, 3: 6}


def test_glossary_is_shared_searchable_and_methods_are_packaged():
    """The same definitions serve app tooltips and downloadable glossary rows."""
    assert len(glossary_rows()) >= 50
    assert field_help(term="Centroid distance") == next(
        row["definition"]
        for row in glossary_rows()
        if row["term"] == "Centroid distance"
    )
    assert "TM-score" in methods_text()
    with pytest.raises(WorkbenchError):
        viewer_asset(name="../outside.js")


def test_safe_paths_hashes_and_json_tsv_helpers(tmp_path):
    """Core serialisation and path checks retain deterministic safe values."""
    path = tmp_path / "a.json"
    write_json(path=path, value={"value": 2})
    assert len(digest_file(path=path)) == 64
    assert "+00:00" in utc_now()
    assert safe_id(value="safe_model-1") == "safe_model-1"
    for relative in ["../escape", "/absolute", "a\\b"]:
        with pytest.raises(WorkbenchError):
            inside(root=tmp_path, relative=relative)
    with pytest.raises(WorkbenchError):
        safe_id(value="unsafe/id")
    write_tsv(path=tmp_path / "test.tsv", rows=[{"a": [1, 2]}], columns=["a"])
    assert "[1, 2]" in (tmp_path / "test.tsv").read_text()
