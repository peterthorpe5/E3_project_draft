"""Batch publication, exact pocket annotation and external-viewer hand-off."""

import copy
import io
import json
import os
import zipfile
from pathlib import Path

import gemmi
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from pocket_workbench import cli, ui
from pocket_workbench.annotations import (
    pocket_alignment_data,
    pocket_alignment_document,
)
from pocket_workbench.building import build_resource
from pocket_workbench.comparison import compare_batch
from pocket_workbench.core import ComparisonSettings, WorkbenchError, load_project
from pocket_workbench.evidence import aligner_agreement, comparison_audit
from pocket_workbench.groups import inspector_bundle, source_link
from pocket_workbench.reporting import export_review
from pocket_workbench.runtime import available_cpus, execution_info


def test_annotations_follow_exact_fasta_columns_and_retain_unmapped(project):
    """Alignment gaps and missing coordinate residues never become guessed sites."""
    data = load_project(root=project)
    aligned = {k: "-" + v for k, v in data["sequences"].items()}
    data["pockets"][0]["residues"].append("A_99999")
    payload = pocket_alignment_data(root=project, data=data, aligned_records=aligned)
    first = next(
        a
        for a in payload["annotations"]
        if a["sample_id"] == "synthetic_reference" and a["residue_key"] == "A_2"
    )
    assert first["sequence_position"] == 2
    assert first["alignment_column"] == 3
    assert (
        next(a for a in payload["annotations"] if a["residue_key"] == "A_99999")[
            "mapping_status"
        ]
        == "unmapped"
    )
    assert payload["rows"][0]["mapped"] < payload["rows"][0]["expected"]
    data["sequences"]["synthetic_reference"] = "WRONG"
    with pytest.raises(WorkbenchError, match="preserve"):
        pocket_alignment_data(root=project, data=data, aligned_records=aligned)


def test_annotation_page_escapes_identifiers_and_preserves_gaps(project):
    """Only generated scripts execute; user data stays in escaped inert JSON."""
    data = load_project(root=project)
    payload = pocket_alignment_data(
        root=project, data=data, aligned_records=data["sequences"]
    )
    payload["rows"][0]["sample_id"] = "</script><script>bad()</script>"
    document = pocket_alignment_document(payload=payload)
    assert "</script><script>bad()" not in document
    assert "Pocket-annotated FASTA alignment" in document
    assert '<script src="http' not in document


def test_local_inspector_bundle_retains_models_and_declares_missing_metrics(project):
    """The E3 hand-off convention carries real mmCIFs and a complete index audit."""
    data = load_project(root=project)
    with zipfile.ZipFile(
        io.BytesIO(inspector_bundle(root=project, models=data["models"]))
    ) as archive:
        names = archive.namelist()
        assert len([n for n in names if n.endswith(".cif")]) == 3
        cif = gemmi.cif.read_string(archive.read("workbench_model_0.cif").decode())
        structure = gemmi.make_structure_from_block(cif.sole_block())
        assert structure[0].count_atom_sites() == 128
        assert b"not_supplied" in archive.read("manifest.tsv")
        assert b"Missing values are unknown" in archive.read("README.txt")
    assert source_link(value="javascript:alert(1)") == ""
    assert source_link(value="https://example.org/model") == "https://example.org/model"
    with pytest.raises(WorkbenchError):
        inspector_bundle(root=project, models=[])


def test_parallel_fit_results_and_exact_cache_are_reused(project, monkeypatch):
    """Parallel fits have one parent database write and reuse validated results."""
    settings = ComparisonSettings(engine="sequence_kabsch")
    pairs = [
        ("synthetic_reference", "synthetic_mobile"),
        ("synthetic_reference", "synthetic_variant"),
    ]
    first = compare_batch(root=project, pairs=pairs, settings=settings, workers=2)
    assert len(first) == 2 and all(r["status"] == "complete" for r in first)
    assert first[0]["pocket_matches"][0]["residue_correspondences"]
    assert first[0]["pocket_matches"][0]["chemical_group_conservation"] == 1
    monkeypatch.setattr(
        "pocket_workbench.comparison.kabsch",
        lambda **_: pytest.fail("A matching completed fit should be reused"),
    )
    second = compare_batch(root=project, pairs=pairs, settings=settings, workers=2)
    assert [r["comparison_id"] for r in first] == [r["comparison_id"] for r in second]
    with pytest.raises(WorkbenchError):
        ComparisonSettings(engine="sequence_kabsch", fit_scope="global")


def test_aligner_agreement_requires_the_same_member_pocket(project):
    """Separate aligners supporting different alternatives cannot create agreement."""
    audit = comparison_audit(root=project, data=load_project(root=project))
    row = dict(
        audit[0],
        engine="usalign",
        rank_one_position_and_fold=False,
        rank_one_conserved=False,
        top_k_tested_pairs=3,
        top_k_fold_supported_member_pockets=[2],
    )
    other = dict(row, engine="tmalign", top_k_fold_supported_member_pockets=[3])
    assert aligner_agreement(rows=[row, other])[0]["top_k_both_aligners"] is False
    other["top_k_fold_supported_member_pockets"] = [2]
    assert aligner_agreement(rows=[row, other])[0]["top_k_both_aligners"] is True
    assert aligner_agreement(rows=[row])[0]["rank_one_both_aligners"] is None


def test_portable_export_contains_offline_review_and_reopenable_resource(
    project, tmp_path
):
    """The published ZIP has HTML, numeric Excel, DuckDB and an exact app resource."""
    out = export_review(root=project, output=tmp_path / "review")
    assert (out / "index.html").is_file()
    assert (out / "pocket_workbench.duckdb").is_file()
    assert list((out / "models").glob("*.html"))
    assert list((out / "alignments").glob("*/pocket_alignment.html"))
    assert list((out / "model_groups").glob("*.zip"))
    workbook = load_workbook(out / "tables/pocket_workbench.xlsx", read_only=True)
    assert "pocket_matches" in workbook.sheetnames
    workbook.close()
    with zipfile.ZipFile(out / "app_resource.zip") as archive:
        assert any(n.endswith("/project.json") for n in archive.namelist())
    with zipfile.ZipFile(out.with_suffix(".zip")) as archive:
        assert "review/index.html" in archive.namelist()
    with pytest.raises(WorkbenchError, match="new review"):
        export_review(root=project, output=out)


@pytest.mark.parametrize("page", ["Structural evidence", "HOG / model groups"])
def test_new_pages_render_precomputed_evidence(project, monkeypatch, tmp_path, page):
    """The generic evidence and HOG source hand-off pages need no compute service."""
    monkeypatch.setenv("POCKET_WORKBENCH_RESOURCE", str(project))
    monkeypatch.setenv("POCKET_WORKBENCH_WORKSPACE", str(tmp_path / "workspace"))
    app = AppTest.from_file(str(Path(ui.__file__)), default_timeout=20).run()
    app.sidebar.radio[0].set_value(page).run()
    assert not app.exception
    assert not app.error


def test_app_browses_by_default_and_opt_in_is_explicit(project, monkeypatch, tmp_path):
    """Opening an app resource does not expose background analysis submit buttons."""
    monkeypatch.setenv("POCKET_WORKBENCH_RESOURCE", str(project))
    monkeypatch.setenv("POCKET_WORKBENCH_WORKSPACE", str(tmp_path / "space"))
    monkeypatch.delenv("POCKET_WORKBENCH_ALLOW_JOBS", raising=False)
    app = AppTest.from_file(str(Path(ui.__file__)), default_timeout=20).run()
    app.sidebar.radio[0].set_value("Compare structures").run()
    assert not any(b.label == "Start structure comparisons" for b in app.button)
    monkeypatch.setenv("POCKET_WORKBENCH_ALLOW_JOBS", "1")
    app.run()
    assert any(b.label == "Start structure comparisons" for b in app.button)
    assert not app.exception


def test_build_reuses_pockets_and_publishes_without_app_jobs(project, tmp_path):
    """A batch build works on prepared evidence with explicit paired Kabsch."""
    original = copy.deepcopy(load_project(root=project)["pockets"])
    assert (
        build_resource(
            resource=project,
            sources=None,
            output=None,
            review_output=tmp_path / "batch_review",
            aligners=["sequence_kabsch"],
            settings=ComparisonSettings(engine="sequence_kabsch"),
            workers=1,
            threads=1,
            bootstrap=0,
        )
        == 0
    )
    assert load_project(root=project)["pockets"] == original
    state = json.loads((project / "provenance/build_status.json").read_text())
    assert state["stage"] == "finished"
    assert state["execution"]["hostname"]
    assert (tmp_path / "batch_review.zip").is_file()


def test_runtime_cli_and_missing_aligner_fail_before_mutation(project, tmp_path):
    """Diagnostics identify the host; missing scientific binaries fail early."""
    assert available_cpus() >= 1
    assert execution_info(root=project)["project_directory"] == str(project)
    assert cli.main(["doctor", "--resource", str(project)]) == 0
    with pytest.raises(WorkbenchError, match="Missing"):
        build_resource(
            resource=project,
            sources=None,
            output=None,
            review_output=tmp_path / "no_review",
            usalign="definitely_no_such_usalign_binary",
            workers=1,
            threads=1,
        )
    assert not (tmp_path / "no_review").exists()


@pytest.mark.parametrize("engine", ["usalign", "tmalign"])
def test_real_global_aligner_retains_unequal_trace_lengths(project, engine):
    """Real aligners see independent traces and recover the known rigid transform."""
    from pocket_workbench.comparison import compare_pair

    executable = os.getenv("POCKET_WORKBENCH_TEST_" + engine.upper())
    if not executable:
        pytest.skip(
            "Set the optional real structural-aligner executable for integration"
        )
    data = load_project(root=project)
    model = next(m for m in data["models"] if m["sample_id"] == "synthetic_mobile")
    pdb = project / model["prepared_model"]
    pdb.write_text(
        "\n".join(
            line
            for line in pdb.read_text().splitlines()
            if not line.startswith("ATOM") or int(line[22:26]) <= 24
        )
        + "\n"
    )
    coordinates = project / model["coordinates"]
    residues = json.loads(coordinates.read_text())[:24]
    coordinates.write_text(json.dumps(residues))
    model["sequence"] = model["sequence"][:24]
    result = compare_pair(
        root=project,
        data=data,
        reference="synthetic_reference",
        mobile="synthetic_mobile",
        settings=ComparisonSettings(engine=engine, fit_scope="global", min_plddt=0),
        usalign=executable,
        tmalign=executable,
    )
    assert result["reference_fit_residues"] == 32
    assert result["mobile_fit_residues"] == 24
    assert result["aligned_atoms"] == 24
    assert result["rmsd"] < 0.01
    assert result["tm_reference"] < result["tm_mobile"]
    assert result["scope"] == "full_model"
    assert "unreported" not in result["tool_version"]
