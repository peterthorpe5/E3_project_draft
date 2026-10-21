"""Portable Python Excel reports following the parent repository's conventions."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .schema import DEFINITIONS, SCHEMAS, TABLE_FIELDS

SHEETS = {
    "samples": "Model_checks",
    "ranked_pockets": "Ranked_cavities",
    "pockets": "All_pockets",
    "residues": "Residue_map",
    "pocket_residues": "Pocket_residues",
    "matches": "Prediction_overlap",
    "method_runs": "Method_runs",
    "fpocket_descriptors": "FPocket_descriptors",
    "run_info": "Run_provenance",
}


def safe_excel_value(*, value: object) -> object:
    """Remove invalid XML control characters while retaining numeric types.

    Args:
        value: Table scalar.

    Returns:
        Excel-safe scalar; very long text explicitly points to the full TSV.
    """
    if isinstance(value, str):
        cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if len(cleaned) > 32767:
            cleaned = cleaned[:32690] + " [Text shortened for Excel; full value in TSV]"
        return cleaned
    return value


def add_sheet(
    *, workbook: Workbook, name: str, columns: list[str], rows: list[dict]
) -> None:
    """Add a typed, filtered, frozen and consistently formatted worksheet.

    Args:
        workbook: Destination workbook.
        name: Unique Excel sheet name.
        columns: Ordered stable field names.
        rows: Numeric/string/null data records.
    """
    if len(rows) > 1_048_575:
        raise ValueError(
            f"{name} exceeds Excel's row limit; query the DuckDB or TSV instead"
        )
    sheet = workbook.create_sheet(title=name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D2" if name == "Ranked_cavities" else "B2"
    sheet.append(iterable=columns)
    for row in rows:
        sheet.append(
            iterable=[safe_excel_value(value=row.get(column)) for column in columns]
        )
    sheet.row_dimensions[1].height = 48
    for cell in sheet[1]:
        cell.fill = PatternFill(fill_type="solid", fgColor="17365D")
        cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(col_idx=index)
        lengths = [len(column)] + [
            len(str(row[column])) for row in rows[:500] if row.get(column) is not None
        ]
        width = min(46, max(15, max(lengths) + 2))
        if column in {"error", "warnings", "description", "value", "residue_ids"}:
            width = 58
        elif column in {
            "sample_id",
            "confidence_status",
            "ranking_basis",
            "sequence_status",
        }:
            width = 30
        sheet.column_dimensions[letter].width = width
        fmt = "General"
        if column.endswith(("_rank", "_count", "_number", "_id", "_length")):
            fmt = "#,##0"
        elif "fraction" in column or column in {
            "fasta_coverage",
            "jaccard",
            "match_jaccard",
        }:
            fmt = "0.0%"
        elif "plddt" in column or column == "volume_a3":
            fmt = "0.0"
        elif "score" in column or "probability" in column:
            fmt = "0.000"
        for cells in sheet.iter_cols(
            min_col=index, max_col=index, min_row=2, max_row=len(rows) + 1
        ):
            for cell in cells:
                cell.number_format = fmt
                cell.font = Font(name="Calibri", size=11, color="17365D")
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=column in {"error", "warnings", "description"},
                )
                if column in {"error", "warnings", "description"} and cell.value:
                    sheet.row_dimensions[cell.row].height = 32
                if isinstance(cell.value, str):
                    # Input labels must remain text even when they begin '='.
                    cell.data_type = "s"
                if column in {"status", "confidence_status"}:
                    if cell.value in {"failed", "partial", "unmapped_residues"}:
                        cell.fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
                    elif cell.value in {"low_local_confidence", "not_assessed"}:
                        cell.fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        if "plddt" in column and rows:
            sheet.conditional_formatting.add(
                range_string=f"{letter}2:{letter}{len(rows) + 1}",
                cfRule=ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="FFF2CC",
                    end_type="num",
                    end_value=100,
                    end_color="DDEBF7",
                ),
            )
    reference = f"A1:{get_column_letter(col_idx=len(columns))}{max(1, len(rows) + 1)}"
    sheet.auto_filter.ref = reference
    if rows:
        table = Table(displayName=f"Pocket_{name}", ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        sheet.add_table(table=table)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def create_workbook(*, path: Path, tables: dict, manifest: dict) -> None:
    """Write the aggregate or one-model workbook with a guide and dictionary.

    Args:
        path: XLSX destination.
        tables: Same records used by TSV, HTML and DuckDB exports.
        manifest: Run provenance and explicit threshold choices.
    """
    workbook = Workbook()
    workbook.remove(worksheet=workbook.active)
    settings = manifest["settings"]
    guide = [
        {
            "topic": "Report",
            "detail": "Predicted pockets from supplied local protein models",
        },
        {"topic": "Run status", "detail": manifest["status"]},
        {
            "topic": "Start here",
            "detail": "Ranked_cavities: fpocket cavities ranked within each model",
        },
        {
            "topic": "Ranking",
            "detail": (
                "P2Rank rescoring rank; fpocket druggability fallback "
                "is explicitly labelled"
            ),
        },
        {
            "topic": "Independent predictions",
            "detail": (
                "All_pockets: filter method=p2rank for independently predicted sites"
            ),
        },
        {"topic": "Confidence interpretation", "detail": settings["confidence"]},
        {
            "topic": "Local confidence rule",
            "detail": "Not assessed (--confidence none)"
            if settings["confidence"] == "none"
            else (
                f"At least {100 * settings['pocket_fraction']:.0f}% of ALL "
                "predicted lining residues have pLDDT >= "
                f"{settings['high_confidence']:g}, with complete mapping"
            ),
        },
        {
            "topic": "Missing values",
            "detail": (
                "Blank means unavailable/not assessed; zero is a measured numeric zero"
            ),
        },
        {
            "topic": "Interpretation",
            "detail": (
                "These are candidate binding sites. Scores do not establish "
                "binding, affinity or therapeutic efficacy"
            ),
        },
        {
            "topic": "Rescoring profile",
            "detail": (
                f"{settings['rescore_config']}; "
                "rescore_2024 is an experimental P2Rank profile"
            ),
        },
        {
            "topic": "Model preparation",
            "detail": (
                "Protein heavy atoms only. No confidence trimming, docking, "
                "PAE analysis or model alignment"
            ),
        },
        {
            "topic": "Original numbering",
            "detail": (
                "Residue_map and Pocket_residues retain author chain, residue, "
                "insertion code and optional FASTA position"
            ),
        },
        {
            "topic": "Sources",
            "detail": "fpocket: https://doi.org/10.1186/1471-2105-10-168",
        },
        {
            "topic": "Sources",
            "detail": "P2Rank: https://doi.org/10.1186/s13321-018-0285-8",
        },
        {
            "topic": "Sources",
            "detail": "fpocket druggability: https://doi.org/10.1021/jm100574m",
        },
        {
            "topic": "Navigation",
            "detail": (
                "Model_checks records every input, including failures. "
                "Method_runs explains tool-specific failures"
            ),
        },
    ]
    add_sheet(workbook=workbook, name="README", columns=["topic", "detail"], rows=guide)
    workbook["README"].column_dimensions["A"].width = 30
    workbook["README"].column_dimensions["B"].width = 105
    for row in workbook["README"].iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        workbook["README"].row_dimensions[row[0].row].height = 32
    for name in (
        "ranked_pockets",
        "samples",
        "pockets",
        "matches",
        "method_runs",
        "residues",
        "pocket_residues",
        "fpocket_descriptors",
        "run_info",
    ):
        add_sheet(
            workbook=workbook,
            name=SHEETS[name],
            columns=TABLE_FIELDS[name],
            rows=tables[name],
        )
    workbook["Run_provenance"].column_dimensions["A"].width = 30
    workbook["Run_provenance"].column_dimensions["B"].width = 105
    for row in workbook["Run_provenance"].iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        workbook["Run_provenance"].row_dimensions[row[0].row].height = 44
    dictionary = [
        {
            "table": table,
            "column": column,
            "type": kind,
            "description": "Run metadata value as JSON text"
            if table == "run_info" and column == "value"
            else DEFINITIONS[column][1],
        }
        for table, schema in SCHEMAS.items()
        for column, kind in schema
    ]
    add_sheet(
        workbook=workbook,
        name="Field_dictionary",
        columns=["table", "column", "type", "description"],
        rows=dictionary,
    )
    workbook["Field_dictionary"].column_dimensions["D"].width = 95
    workbook.properties.creator = "Protein Pocket Pipeline"
    workbook.properties.title = "Protein pocket evidence"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    workbook.save(filename=temporary)
    temporary.replace(target=path)
