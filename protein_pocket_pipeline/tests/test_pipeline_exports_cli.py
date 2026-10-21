"""End-to-end contracts, failure recovery, format consistency and the CLI."""

from __future__ import annotations

import hashlib
import json
import runpy
import sys
from dataclasses import asdict
from pathlib import Path
from unittest.mock import Mock

import duckdb
import pytest
from conftest import emulate_command
from openpyxl import load_workbook

from protein_pockets.cli import build_parser, main
from protein_pockets.excel import create_workbook, safe_excel_value
from protein_pockets.io import sha256_file
from protein_pockets.models import Sample, Settings
from protein_pockets.pipeline import (
    analyse_method,
    analyse_sample,
    checkpoint_valid,
    run_pipeline,
)
from protein_pockets.reporting import create_database, create_html
from protein_pockets.schema import SCHEMAS, TABLE_FIELDS
from protein_pockets.tools import inspect_tools, resolve_executable, run_command


def test_complete_run_and_consistent_outputs(
    *, tmp_path: Path, model_path: Path, emulated_tools: dict
) -> None:
    """One execution produces consistent ranked Excel, TSV, HTML and DuckDB."""
    root = tmp_path / "results with spaces"
    code = run_pipeline(
        samples=[Sample(sample_id="model_1", model_path=model_path, sequence="ACD")],
        output=root,
        tools=emulated_tools,
        settings=Settings(),
        resume=False,
    )
    assert code == 0
    assert not list(root.rglob("*.csv"))
    with duckdb.connect(
        database=str(root / "database/protein_pockets.duckdb"), read_only=True
    ) as database:
        assert database.execute(
            query="SELECT pocket_id FROM ranked_pockets ORDER BY review_rank"
        ).fetchall() == [(2,), (1,)]
        assert database.execute(query="SELECT COUNT(*) FROM pockets").fetchone()[0] == 5
        assert (
            database.execute(
                query="SELECT COUNT(*) FROM v_best_pocket_per_model"
            ).fetchone()[0]
            == 1
        )
        assert (
            database.execute(
                query="SELECT COUNT(*) FROM v_independent_p2rank"
            ).fetchone()[0]
            == 1
        )
        assert (
            database.execute(
                query="SELECT COUNT(*) FROM v_locally_confident_pockets"
            ).fetchone()[0]
            == 0
        )
        assert (
            database.execute(
                query="SELECT cavity_count FROM v_model_summary"
            ).fetchone()[0]
            == 2
        )
        assert (
            database.execute(
                query="SELECT value FROM fpocket_descriptors "
                "WHERE descriptor='total_sasa'"
            ).fetchone()[0]
            == 123.4
        )
    workbook = load_workbook(
        filename=root / "excel/protein_pockets.xlsx", data_only=False
    )
    ranked = workbook["Ranked_cavities"]
    assert ranked.freeze_panes == "D2"
    assert len(ranked.tables) == 1
    assert ranked["C2"].value == 2
    assert ranked["F2"].data_type == "n"
    assert ranked["E2"].number_format == "#,##0"
    assert ranked["N2"].value == "low_local_confidence"
    assert len(workbook.sheetnames) == 11
    assert workbook["Field_dictionary"].max_row == 1 + sum(
        len(schema) for schema in SCHEMAS.values()
    )
    assert (root / "excel/models/model_1.xlsx").is_file()
    html = (root / "report.html").read_text(encoding="utf-8")
    assert "__REPORT_DATA__" not in html
    assert 'src="http' not in html
    assert "model_1" in html
    assert not (root / ".run.lock").exists()
    assert (
        json.loads((root / "provenance/run_manifest.json").read_text(encoding="utf-8"))[
            "status"
        ]
        == "complete"
    )


def test_resume_verifies_checksums_and_refuses_changed_settings(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Intact checkpoints skip tools; altered files trigger real model reanalysis."""
    root = tmp_path / "results"
    options = {
        "samples": [Sample(sample_id="model", model_path=model_path)],
        "output": root,
        "tools": emulated_tools,
        "settings": Settings(),
    }
    assert run_pipeline(**options, resume=False) == 0
    runner = Mock(side_effect=emulate_command)
    monkeypatch.setattr("protein_pockets.pipeline.run_command", runner)
    assert run_pipeline(**options, resume=True) == 0
    assert runner.call_count == 0
    prepared = root / "models/model/prepared.pdb"
    prepared.write_text(data="corrupted", encoding="utf-8")
    assert run_pipeline(**options, resume=True) == 0
    assert runner.call_count == 3
    assert "corrupted" not in prepared.read_text(encoding="utf-8")
    with pytest.raises(expected_exception=ValueError, match="Output is not empty"):
        run_pipeline(**options, resume=False)
    options["settings"] = Settings(threads=2)
    with pytest.raises(expected_exception=ValueError, match="Resume refused"):
        run_pipeline(**options, resume=True)


def test_failed_method_keeps_independent_results(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Rescoring failure leaves fpocket and independent P2Rank usable and labelled."""

    def fail_rescore(*, command: list[str], **options: object) -> dict:
        """Fail only the rescoring stage in this behavioural test."""
        if "rescore" in command:
            raise RuntimeError("Deliberate unit-test tool failure")
        return emulate_command(command=command, **options)

    monkeypatch.setattr("protein_pockets.pipeline.run_command", fail_rescore)
    root = tmp_path / "partial"
    code = run_pipeline(
        samples=[Sample(sample_id="model", model_path=model_path)],
        output=root,
        tools=emulated_tools,
        settings=Settings(),
        resume=False,
    )
    assert code == 2
    with duckdb.connect(
        database=str(root / "database/protein_pockets.duckdb"), read_only=True
    ) as database:
        assert (
            database.execute(query="SELECT status FROM samples").fetchone()[0]
            == "partial"
        )
        assert (
            database.execute(
                query="SELECT DISTINCT ranking_basis FROM ranked_pockets"
            ).fetchone()[0]
            == "fpocket_druggability"
        )
        assert (
            database.execute(
                query="SELECT p2rank_rescore_score FROM ranked_pockets LIMIT 1"
            ).fetchone()[0]
            is None
        )
    assert (root / "report.html").exists()


def test_preparation_failure_is_accounted_for(
    *, tmp_path: Path, model_path: Path, emulated_tools: dict
) -> None:
    """A sequence mismatch remains visible alongside completed models."""
    root = tmp_path / "mixed"
    code = run_pipeline(
        samples=[
            Sample(sample_id="valid", model_path=model_path),
            Sample(sample_id="invalid", model_path=model_path, sequence="WWW"),
        ],
        output=root,
        tools=emulated_tools,
        settings=Settings(),
        resume=False,
    )
    assert code == 2
    with duckdb.connect(
        database=str(root / "database/protein_pockets.duckdb"), read_only=True
    ) as database:
        assert database.execute(query="SELECT COUNT(*) FROM samples").fetchone()[0] == 2
        assert (
            database.execute(
                query="SELECT residue_count FROM samples WHERE sample_id='invalid'"
            ).fetchone()[0]
            is None
        )
        assert (
            database.execute(
                query="SELECT COUNT(*) FROM method_runs WHERE status='skipped'"
            ).fetchone()[0]
            == 3
        )
    assert (root / "excel/models/invalid.xlsx").exists()


def test_empty_results_and_typed_database(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """No pockets is a valid, reportable outcome with typed empty tables."""
    monkeypatch.setattr("protein_pockets.pipeline.analyse_method", lambda **_: [])
    root = tmp_path / "no_pockets"
    assert (
        run_pipeline(
            samples=[Sample(sample_id="model", model_path=model_path)],
            output=root,
            tools=emulated_tools,
            settings=Settings(),
            resume=False,
        )
        == 0
    )
    with duckdb.connect(
        database=str(root / "database/protein_pockets.duckdb"), read_only=True
    ) as database:
        assert (
            database.execute(query="SELECT status FROM samples").fetchone()[0]
            == "no_pockets"
        )
        assert (
            database.execute(query="SELECT COUNT(*) FROM ranked_pockets").fetchone()[0]
            == 0
        )
        types = database.execute(query="DESCRIBE ranked_pockets").fetchall()
        assert any(row[0] == "volume_a3" and row[1] == "DOUBLE" for row in types)
    workbook = load_workbook(filename=root / "excel/protein_pockets.xlsx")
    assert workbook["Ranked_cavities"].max_row == 1


def test_every_tool_can_fail_without_losing_model_record(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """All-tool failure gives exit two and a retained model-level error."""
    monkeypatch.setattr(
        "protein_pockets.pipeline.analyse_method",
        Mock(side_effect=RuntimeError("failed")),
    )
    tables = analyse_sample(
        sample=Sample(sample_id="model", model_path=model_path),
        root=tmp_path / "results",
        tools=emulated_tools,
        settings=Settings(),
    )
    assert tables["samples"][0]["status"] == "failed"
    assert len(tables["method_runs"]) == 3
    assert len(tables["pockets"]) == 0


def test_disabled_rescoring(
    *, tmp_path: Path, model_path: Path, emulated_tools: dict
) -> None:
    """Deliberately skipping rescoring is successful and changes the ranking basis."""
    tables = analyse_sample(
        sample=Sample(sample_id="model", model_path=model_path),
        root=tmp_path / "results",
        tools=emulated_tools,
        settings=Settings(rescore=False),
    )
    assert tables["samples"][0]["status"] == "complete"
    assert tables["samples"][0]["rescore_status"] == "skipped"
    assert tables["ranked_pockets"][0]["ranking_basis"] == "fpocket_druggability"


def test_later_stage_cannot_leave_an_earlier_native_csv(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Late tool writes are converted after the complete model analysis."""

    def recreate_prediction(*, command: list[str], **options: object) -> dict:
        """Emulate a later stage recreating an earlier native output table."""
        result = emulate_command(command=command, **options)
        if "rescore" in command:
            path = options["cwd"].parent / "p2rank/output/model.pdb_predictions.csv"
            path.write_text(data="name,rank\npocket1,1\n", encoding="utf-8")
        return result

    monkeypatch.setattr("protein_pockets.pipeline.run_command", recreate_prediction)
    root = tmp_path / "late_native_write"
    tables = analyse_sample(
        sample=Sample(sample_id="model", model_path=model_path),
        root=root,
        tools=emulated_tools,
        settings=Settings(),
    )
    assert tables["samples"][0]["status"] == "complete"
    assert not list(root.rglob("*.csv"))


def test_checkpoint_safety_and_changed_files(*, tmp_path: Path) -> None:
    """Cache paths stay within the root; changed or missing files invalidate them."""
    path = tmp_path / "file.tsv"
    path.write_text(data="correct", encoding="utf-8")
    tables = {name: [] for name in TABLE_FIELDS}
    tables["samples"] = [{"sample_id": "model", "status": "complete"}]
    checkpoint = {
        "tables": tables,
        "tables_sha256": hashlib.sha256(
            json.dumps(tables, sort_keys=True).encode()
        ).hexdigest(),
        "checksums": {"file.tsv": sha256_file(path=path)},
    }
    assert checkpoint_valid(checkpoint=checkpoint, root=tmp_path)
    assert not checkpoint_valid(checkpoint=checkpoint, root=tmp_path, sample_id="other")
    assert not checkpoint_valid(checkpoint=[], root=tmp_path)
    assert not checkpoint_valid(checkpoint={"tables": []}, root=tmp_path)
    path.unlink()
    assert not checkpoint_valid(checkpoint=checkpoint, root=tmp_path)
    checkpoint["checksums"] = {"../outside": "any"}
    assert not checkpoint_valid(checkpoint=checkpoint, root=tmp_path)
    checkpoint["checksums"] = {}
    assert not checkpoint_valid(checkpoint=checkpoint, root=tmp_path)
    checkpoint["tables"]["samples"][0]["status"] = "failed"
    assert not checkpoint_valid(checkpoint=checkpoint, root=tmp_path)


def test_unrecognised_output_and_internal_inputs_are_refused(
    *, tmp_path: Path, model_path: Path, tool_identity: dict
) -> None:
    """The pipeline never adopts unrelated output folders or inputs inside them."""
    with pytest.raises(expected_exception=ValueError, match="outside"):
        run_pipeline(
            samples=[Sample(sample_id="model", model_path=model_path)],
            output=tmp_path,
            tools=tool_identity,
            settings=Settings(),
            resume=False,
        )
    root = tmp_path / "unrelated"
    root.mkdir()
    (root / "unrelated.txt").touch()
    with pytest.raises(expected_exception=ValueError, match="Output is not empty"):
        run_pipeline(
            samples=[Sample(sample_id="model", model_path=model_path)],
            output=root,
            tools=tool_identity,
            settings=Settings(),
            resume=True,
        )
    with pytest.raises(expected_exception=ValueError, match="At least one"):
        run_pipeline(
            samples=[],
            output=root,
            tools=tool_identity,
            settings=Settings(),
            resume=False,
        )


def test_lock_and_export_failure_cleanup(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Concurrent runs are rejected and export failures retain resumable state."""
    root = tmp_path / "results"
    options = {
        "samples": [Sample(sample_id="model", model_path=model_path)],
        "output": root,
        "tools": emulated_tools,
        "settings": Settings(),
    }
    assert run_pipeline(**options, resume=False) == 0
    (root / ".run.lock").touch()
    with pytest.raises(expected_exception=ValueError, match="locked"):
        run_pipeline(**options, resume=True)
    (root / ".run.lock").unlink()
    monkeypatch.setattr(
        "protein_pockets.reporting.export_results",
        Mock(side_effect=OSError("disk full")),
    )
    with pytest.raises(expected_exception=OSError, match="disk full"):
        run_pipeline(**options, resume=True)
    assert not (root / ".run.lock").exists()
    assert (
        json.loads((root / "provenance/run_manifest.json").read_text(encoding="utf-8"))[
            "status"
        ]
        == "interrupted_or_export_failed"
    )


def test_corrupt_checkpoint_retries(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A damaged JSON checkpoint is retried rather than breaking all results."""
    root = tmp_path / "results"
    options = {
        "samples": [Sample(sample_id="model", model_path=model_path)],
        "output": root,
        "tools": emulated_tools,
        "settings": Settings(),
    }
    assert run_pipeline(**options, resume=False) == 0
    (root / "checkpoints/model.json").write_text(data="broken", encoding="utf-8")
    runner = Mock(side_effect=emulate_command)
    monkeypatch.setattr("protein_pockets.pipeline.run_command", runner)
    assert run_pipeline(**options, resume=True) == 0
    assert runner.call_count == 3


def test_database_has_no_external_dependencies(*, tmp_path: Path) -> None:
    """Empty databases remain queryable and enforce unique record identities."""
    tables = {name: [] for name in TABLE_FIELDS}
    path = tmp_path / "pockets.duckdb"
    create_database(path=path, tables=tables)
    with duckdb.connect(database=str(path), read_only=True) as database:
        assert (
            database.execute(
                query="SELECT COUNT(*) FROM v_best_pocket_per_model"
            ).fetchone()[0]
            == 0
        )
    tables["samples"] = [{"sample_id": "same"}, {"sample_id": "same"}]
    with pytest.raises(expected_exception=duckdb.ConstraintException):
        create_database(path=path, tables=tables)
    with duckdb.connect(database=str(path), read_only=True) as database:
        assert database.execute(query="SELECT COUNT(*) FROM samples").fetchone()[0] == 0


def test_excel_text_is_safe_and_numeric_types_are_retained(*, tmp_path: Path) -> None:
    """Formula-like metadata remains text; numbers and missing values stay typed."""
    tables = {name: [] for name in TABLE_FIELDS}
    tables["samples"] = [
        {
            "sample_id": "model",
            "warnings": "=1+1",
            "residue_count": 3,
            "mean_plddt": None,
        }
    ]
    path = tmp_path / "safe.xlsx"
    create_workbook(
        path=path,
        tables=tables,
        manifest={"settings": asdict(Settings()), "status": "complete"},
    )
    workbook = load_workbook(filename=path, data_only=False)
    sheet = workbook["Model_checks"]
    header = [cell.value for cell in sheet[1]]
    warning = sheet.cell(row=2, column=header.index("warnings") + 1)
    assert warning.value == "=1+1"
    assert warning.data_type == "s"
    assert safe_excel_value(value="x\x00y") == "xy"
    assert "full value in TSV" in safe_excel_value(value="x" * 40000)
    assert safe_excel_value(value=1.2) == 1.2


def test_html_escapes_script_injection(*, tmp_path: Path) -> None:
    """Metadata containing HTML cannot escape the JSON data block."""
    path = tmp_path / "report.html"
    tables = {name: [] for name in TABLE_FIELDS}
    tables["samples"] = [
        {
            "sample_id": "model",
            "warnings": '</script><script>alert("bad")</script>&\u2028',
        }
    ]
    create_html(
        path=path,
        tables=tables,
        manifest={
            "status": "complete",
            "finished_utc": "now",
            "settings": asdict(Settings()),
            "pipeline_version": "test",
        },
    )
    text = path.read_text(encoding="utf-8")
    assert '<script>alert("bad")' not in text
    assert "\\u003c/script\\u003e" in text
    assert "\\u2028" in text


def test_command_success_failure_and_timeout(*, tmp_path: Path) -> None:
    """Real subprocesses record logs and get terminated on timeouts."""
    record = run_command(
        command=[sys.executable, "-c", "print('ok')"],
        cwd=tmp_path,
        log_prefix=tmp_path / "success",
        timeout=5,
    )
    assert record["return_code"] == 0
    assert "ok" in (tmp_path / "success.stdout.log").read_text(encoding="utf-8")
    with pytest.raises(expected_exception=RuntimeError, match="Tool failure"):
        run_command(
            command=[sys.executable, "-c", "raise SystemExit(3)"],
            cwd=tmp_path,
            log_prefix=tmp_path / "failure",
            timeout=5,
        )
    with pytest.raises(expected_exception=RuntimeError, match="Timeout"):
        run_command(
            command=[sys.executable, "-c", "import time; time.sleep(30)"],
            cwd=tmp_path,
            log_prefix=tmp_path / "timeout",
            timeout=0.1,
        )
    assert (
        json.loads((tmp_path / "timeout.command.json").read_text(encoding="utf-8"))[
            "timeout"
        ]
        is True
    )


def test_resolve_tools_and_asset_identity(*, tmp_path: Path) -> None:
    """Tool preflight checks versions and fingerprints selected model weights."""
    executable = tmp_path / "prank"
    executable.write_text(
        data="#!/bin/sh\nprintf 'P2Rank 2.5.1\\n'\n", encoding="utf-8"
    )
    executable.chmod(mode=0o755)
    config = tmp_path / "config"
    config.mkdir()
    (config / "alphafold.groovy").write_text(data="test", encoding="utf-8")
    weights = tmp_path / "models/alphafold"
    weights.mkdir(parents=True)
    (weights / "model.zst").write_bytes(data=b"fixture")
    assert resolve_executable(value=str(executable)) == executable
    record = inspect_tools(
        fpocket=str(executable), p2rank=str(executable), settings=Settings()
    )
    assert "models/alphafold/model.zst" in record["p2rank"]["assets"]
    with pytest.raises(expected_exception=ValueError, match="version check failed"):
        inspect_tools(
            fpocket=str(executable),
            p2rank=str(executable),
            settings=Settings(required_p2rank="different"),
        )
    executable.write_text(data="#!/bin/sh\nexit 0\n", encoding="utf-8")
    with pytest.raises(expected_exception=ValueError, match="no identifiable"):
        inspect_tools(
            fpocket=str(executable), p2rank=str(executable), settings=Settings()
        )
    with pytest.raises(expected_exception=ValueError, match="not found"):
        resolve_executable(value="definitely-not-an-installed-tool-123")


def test_cli_check_and_run(
    *,
    tmp_path: Path,
    model_path: Path,
    emulated_tools: dict,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture,
) -> None:
    """The documented short CLI routes work without configuration files."""
    monkeypatch.setattr("protein_pockets.cli.inspect_tools", lambda **_: emulated_tools)
    assert main(argv=["check"]) == 0
    assert "fixture-prank" in capsys.readouterr().out
    output = tmp_path / "tool_versions.json"
    assert main(argv=["check", "--output", str(output)]) == 0
    assert json.loads(output.read_text(encoding="utf-8")) == emulated_tools
    root = tmp_path / "cli_results"
    assert (
        main(
            argv=[
                "run",
                "--models",
                str(model_path.parent),
                "--output",
                str(root),
                "--skip-rescore",
            ]
        )
        == 0
    )
    assert (
        main(
            argv=[
                "run",
                "--models",
                str(model_path.parent),
                "--output",
                str(tmp_path / "other"),
                "--threads",
                "0",
            ]
        )
        == 1
    )
    assert build_parser().parse_args(args=["check"]).command == "check"


def test_cli_interruption_and_module_entrypoint(
    *, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Ctrl-C returns a useful status and python -m exposes the same CLI."""
    monkeypatch.setattr(
        "protein_pockets.cli.inspect_tools", Mock(side_effect=KeyboardInterrupt)
    )
    assert main(argv=["check"]) == 130
    monkeypatch.setattr(sys, "argv", ["protein_pockets", "--version"])
    with pytest.raises(expected_exception=SystemExit) as caught:
        runpy.run_module(mod_name="protein_pockets", run_name="__main__")
    assert caught.value.code == 0


def test_method_output_contracts(
    *,
    tmp_path: Path,
    model_path: Path,
    tool_identity: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Missing predictions or incomplete rescoring fail with explicit contracts."""

    def empty_command(
        *, command: list[str], cwd: Path, log_prefix: Path, timeout: float
    ) -> dict:
        """Create no native outputs while pretending the process succeeded."""
        log_prefix.parent.mkdir(parents=True, exist_ok=True)
        log_prefix.with_suffix(".stdout.log").write_text(
            data="No pockets found\n", encoding="utf-8"
        )
        log_prefix.with_suffix(".stderr.log").write_text(data="", encoding="utf-8")
        return {}

    monkeypatch.setattr("protein_pockets.pipeline.run_command", empty_command)
    options = {
        "sample": Sample(sample_id="model", model_path=model_path),
        "root": tmp_path / "root",
        "prepared": model_path,
        "tools": tool_identity,
        "settings": Settings(),
        "previous_pockets": [],
    }
    assert analyse_method(**options, method="fpocket") == []
    with pytest.raises(expected_exception=ValueError, match="predictions file"):
        analyse_method(**options, method="p2rank")
    with pytest.raises(expected_exception=ValueError, match="Unknown method"):
        analyse_method(**options, method="invalid")


@pytest.mark.external
def test_real_external_tool_smoke(*, tmp_path: Path) -> None:
    """Optionally run real tools on a caller-supplied, local test structure."""
    import os

    values = {
        key: os.environ.get(key)
        for key in ("POCKET_TEST_PDB", "POCKET_TEST_FPOCKET", "POCKET_TEST_P2RANK")
    }
    if not all(values.values()):
        pytest.skip("Set POCKET_TEST_PDB, POCKET_TEST_FPOCKET and POCKET_TEST_P2RANK")
    settings = Settings(confidence="none", threads=2)
    tools = inspect_tools(
        fpocket=values["POCKET_TEST_FPOCKET"],
        p2rank=values["POCKET_TEST_P2RANK"],
        settings=settings,
    )
    assert (
        run_pipeline(
            samples=[
                Sample(
                    sample_id="external_smoke",
                    model_path=Path(values["POCKET_TEST_PDB"]),
                )
            ],
            output=tmp_path / "real_results",
            tools=tools,
            settings=settings,
            resume=False,
        )
        == 0
    )
    with duckdb.connect(
        database=str(tmp_path / "real_results/database/protein_pockets.duckdb"),
        read_only=True,
    ) as database:
        assert (
            database.execute(query="SELECT COUNT(*) FROM ranked_pockets").fetchone()[0]
            > 0
        )
